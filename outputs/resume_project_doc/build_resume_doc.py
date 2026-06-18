from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Unity_Roguelite_resume_project.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(15, 23, 42)
GRAY = RGBColor(89, 89, 89)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9DEE7", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_tbl_width_fixed(table, widths_in):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para_format(p, before=0, after=6, line=1.10):
    paragraph_format = p.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = line


def add_title(doc):
    p = doc.add_paragraph()
    set_para_format(p, before=0, after=4, line=1.0)
    run = p.add_run("Unity Roguelite 回合制手游项目简历补充材料")
    set_run_font(run, size=23, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    set_para_format(subtitle, before=0, after=10, line=1.15)
    run = subtitle.add_run("适用于游戏开发 / Unity 客户端 / 玩法程序方向简历")
    set_run_font(run, size=11, color=GRAY)

    rows = [
        ("项目类型", "团队合作 Unity 移动端 Roguelite 回合制游戏"),
        ("技术关键词", "Unity/Tuanjie 2022.3、C#、ScriptableObject、UGUI、回合制战斗、数据驱动配置"),
        ("个人职责", "素材收集、地图设计、部分怪物设计、怪物相关代码、遗物相关代码"),
        ("使用方式", "可直接选取“简历精简版”放入简历，并结合“面试展开点”准备说明"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        set_para_format(p, before=0, after=2, line=1.1)
        label_run = p.add_run(label + "：")
        set_run_font(label_run, size=10.5, color=INK, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10.5, color=INK)

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2F3")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    set_para_format(rule, before=6, after=12, line=1.0)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, color=BLUE, bold=True)
        set_para_format(p, before=16, after=8, line=1.1)
    elif level == 2:
        set_run_font(run, size=13, color=BLUE, bold=True)
        set_para_format(p, before=12, after=6, line=1.1)
    else:
        set_run_font(run, size=12, color=DARK_BLUE, bold=True)
        set_para_format(p, before=8, after=4, line=1.1)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_para_format(p, before=0, after=6, line=1.10)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, size=11, color=INK, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run, size=11, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para_format(p, before=0, after=8, line=1.167)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D9E2F3", size="6")
    set_tbl_width_fixed(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    set_para_format(p, before=0, after=2, line=1.12)
    run = p.add_run(title + "：")
    set_run_font(run, size=10.5, color=DARK_BLUE, bold=True)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    spacer = doc.add_paragraph()
    set_para_format(spacer, before=0, after=2, line=1.0)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_tbl_width_fixed(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        set_para_format(p, before=0, after=0, line=1.1)
        run = p.add_run(header)
        set_run_font(run, size=10.5, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            set_para_format(p, before=0, after=0, line=1.1)
            run = p.add_run(value)
            set_run_font(run, size=10.5, color=INK)
    set_tbl_width_fixed(table, widths)
    spacer = doc.add_paragraph()
    set_para_format(spacer, before=0, after=4, line=1.0)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = INK
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = "Unity Roguelite 项目简历补充材料"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_format(header, before=0, after=0, line=1.0)
    for run in header.runs:
        set_run_font(run, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("简历项目经历素材")
    set_run_font(run, size=9, color=GRAY)

    add_title(doc)
    add_heading(doc, "一、项目概述", 1)
    add_body(
        doc,
        "本项目是一款基于 Unity/Tuanjie 2022.3 与 C# 开发的移动端 Roguelite 回合制游戏。"
        "游戏核心流程围绕地图探索、骰子拼点战斗、怪物遭遇、遗物构筑、事件、商店、宝箱与角色成长展开。"
        "项目采用 ScriptableObject 管理敌人、技能、遗物、地图节点和奖励内容，便于团队进行数据驱动配置与玩法迭代。",
    )
    add_callout(
        doc,
        "简历定位",
        "建议将该项目放在“项目经历”中，突出 Unity 客户端开发、玩法系统实现、数据驱动配置和团队协作能力。",
    )

    add_heading(doc, "二、简历精简版", 1)
    add_body(
        doc,
        "Unity Roguelite 回合制手游 | Unity/Tuanjie 2022.3、C#、ScriptableObject",
        bold_prefix="Unity Roguelite 回合制手游",
    )
    for item in [
        "与团队合作开发移动端 Roguelite 回合制游戏，参与地图探索、骰子战斗、怪物、遗物和奖励等核心玩法模块。",
        "负责素材收集与资源整理，完成角色、怪物、地图、UI 等素材的筛选、导入和项目内归档。",
        "参与地图关卡设计，规划战斗、事件、宝箱、商店等节点分布，优化玩家探索路线与单局节奏。",
        "设计并配置部分怪物，包括属性、技能、状态效果和战斗定位，并实现部分怪物技能与战斗结算逻辑。",
        "参与遗物系统开发，实现部分遗物触发条件和效果逻辑，支持遗物与战斗、地图、奖励等系统联动。",
        "基于 ScriptableObject 进行数据驱动配置，降低内容扩展成本，提升怪物、技能和遗物的迭代效率。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "三、职责拆解版", 1)
    add_table(
        doc,
        ["职责方向", "可写入简历的表达"],
        [
            ("素材收集", "负责游戏素材收集、筛选、分类与 Unity 项目导入适配，保证角色、怪物、地图、UI 等资源符合项目风格与使用需求。"),
            ("地图设计", "参与 Roguelite 地图节点设计，规划战斗、事件、商店、宝箱、休息等节点分布，优化探索节奏和风险收益体验。"),
            ("怪物设计", "参与部分怪物机制设计，制定敌人属性、技能效果、状态能力和战斗定位，并完成相关数据配置。"),
            ("怪物代码", "实现部分怪物技能触发、攻击结算、状态施加、目标选择等战斗逻辑，配合战斗系统完成玩法闭环。"),
            ("遗物代码", "参与遗物系统开发，实现部分遗物触发条件、效果结算和跨系统联动逻辑，支撑玩家构筑体验。"),
            ("团队协作", "配合团队完成数值调试、玩法验证和内容迭代，根据测试反馈调整怪物强度、地图节奏和遗物效果。"),
        ],
        [1.35, 5.15],
    )

    add_heading(doc, "四、推荐正式写法", 1)
    add_body(doc, "可直接放入简历的完整表述如下：")
    for item in [
        "负责部分怪物与遗物玩法模块开发，基于 ScriptableObject 配置敌人属性、技能、遗物效果与奖励内容。",
        "实现怪物技能触发、状态施加、目标选择和战斗结算逻辑，提升敌人机制差异化与战斗策略性。",
        "参与遗物触发与效果结算逻辑开发，支持遗物与战斗、地图、奖励、经济等系统产生联动。",
        "参与 Roguelite 地图节点与探索路线设计，规划战斗、事件、宝箱、商店等节点分布，优化单局成长节奏。",
        "负责项目素材收集与资源整理，保证美术、音效和 UI 资源能够顺利接入 Unity 项目并服务玩法需求。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "五、面试可展开亮点", 1)
    for title, text in [
        ("怪物不是单纯填数值", "可以说明自己会根据敌人的定位设计机制，例如输出型、防御型、状态型、干扰型敌人，再把设计转化为技能配置和代码逻辑。"),
        ("遗物系统强调触发时机", "可以说明遗物会在战斗开始、回合开始、造成伤害、受到伤害、获得奖励、移动或商店购买等时机触发，自己参与了其中部分效果实现与系统接入。"),
        ("地图设计关注节奏", "可以说明地图节点分布要考虑风险收益、成长路线、奖励密度和 Boss 前准备空间，而不是随机摆放节点。"),
        ("数据驱动便于迭代", "可以说明项目使用 ScriptableObject 管理怪物、技能、遗物和地图内容，让策划和程序都能更快调整内容。"),
    ]:
        p = doc.add_paragraph()
        set_para_format(p, before=0, after=6, line=1.12)
        run = p.add_run(title + "：")
        set_run_font(run, size=11, color=DARK_BLUE, bold=True)
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)

    add_heading(doc, "六、可补充信息", 1)
    add_body(doc, "如果你希望简历更具体，可以继续补充以下信息，我可以再帮你改成更强的一版：")
    for item in [
        "项目周期：例如 2026.03 - 2026.06。",
        "团队规模：例如 2 人、3 人或 5 人。",
        "你实现过的具体怪物名称或机制。",
        "你实现过的具体遗物名称或效果。",
        "是否完成可运行 Demo、答辩展示或发布测试包。",
    ]:
        add_bullet(doc, item)

    props = doc.core_properties
    props.title = "Unity Roguelite 回合制手游项目简历补充材料"
    props.subject = "简历项目经历"
    props.author = "Codex"
    props.comments = "Generated for resume project experience drafting."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
